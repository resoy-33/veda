"""Deterministic French Excel report (the human-readable deliverable)."""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core.canonical import sha256_bytes

AMOUNT_CODES = {
    "marge_brute",
    "ebitda",
    "total_dettes",
    "dette_nette",
    "fonds_roulement",
    "bfr_exploitation",
    "tresorerie_nette_gestion",
    "capitaux_investis",
    "free_cash_flow",
    "fcff",
    "capitalisation_boursiere",
    "valeur_entreprise",
}
DAY_CODES = {"dio", "dso", "dpo", "ccc"}
PERCENT_CODES = {
    "marge_brute_pct",
    "marge_exploitation",
    "marge_nette",
    "roe",
    "roa",
    "roce",
}
KEY_RATIOS = (
    "marge_brute",
    "ebitda",
    "free_cash_flow",
    "dette_nette",
    "marge_brute_pct",
    "marge_exploitation",
    "marge_nette",
    "roe",
    "roa",
    "roce",
    "liquidite_generale",
    "liquidite_reduite",
    "gearing",
    "dette_nette_ebitda",
    "couverture_interets_ebit",
    "dio",
    "dso",
    "dpo",
    "ccc",
)
FIXED_DATE = datetime(2026, 1, 1, 0, 0, 0)
HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=15, bold=True, color="1F3B57")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CORE_TIMESTAMP = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)


def report_filename(issuer: str, year: int) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", issuer).strip("_") or "depot"
    return f"rapport_{safe}_{year}.xlsx"


def _mtnd(value: float) -> float:
    return value / 1_000_000


def _ratio_format(code: str) -> str:
    if code in AMOUNT_CODES:
        return "#,##0"
    if code in DAY_CODES:
        return "0.0"
    return "0.000"


def _write_table(ws, headers: list[str], rows: list[list], formats: list[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for index, fmt in enumerate(formats, start=1):
        letter = get_column_letter(index)
        for cell in ws[letter][1:]:
            cell.border = BORDER
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _auto_width(ws):
    for column in ws.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(12, width + 2), 62)


def _summary_sheet(
    ws,
    filing: dict,
    ratios: list[dict],
    checks: list[dict],
    extraction_sha256: str,
    ratios_sha256: str,
):
    ws["A1"] = "Rapport de dépôt vérifié"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Veda — couche de confiance pour les états financiers"
    ws["A2"].font = Font(italic=True, color="5A6B7B")
    by_code = {row["code"]: row for row in ratios}
    verdict = "Conforme" if all(c["status"] == "pass" for c in checks) else "Écart"
    rows = [
        ("Émetteur", filing["issuer"], ""),
        ("Période", str(filing["period"])[:10], ""),
        ("Secteur", filing.get("sector", ""), ""),
        ("Devise", filing.get("unit", "TND"), ""),
        ("Consolidation", filing.get("consolidation", "individuel"), ""),
        ("Empreinte du document (SHA-256)", filing.get("doc_sha256", ""), ""),
        ("Empreinte de l'extraction (SHA-256)", extraction_sha256, ""),
        ("Empreinte des ratios (SHA-256)", ratios_sha256, ""),
        ("Verdict des contrôles", verdict, ""),
        ("Topic HCS", filing.get("topic_id") or "—", ""),
        ("Identifiant du dépôt", filing.get("id", ""), ""),
    ]
    start = 4
    for offset, (label, value, _) in enumerate(rows):
        ws.cell(row=start + offset, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=start + offset, column=2, value=value)
        cell.alignment = Alignment(wrap_text=False)
    ratio_start = start + len(rows) + 2
    ws.cell(row=ratio_start - 1, column=1, value="Indicateurs clés").font = Font(
        bold=True, size=12, color="1F3B57"
    )
    for column, header in enumerate(("Indicateur", "Valeur", "Unité"), start=1):
        cell = ws.cell(row=ratio_start, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    row_index = ratio_start + 1
    for code in KEY_RATIOS:
        row = by_code.get(code)
        if row is None:
            continue
        value = row["value"]
        if value is None:
            display, unit = "—", ""
        elif code in AMOUNT_CODES:
            display, unit = _mtnd(value), "MTND"
        elif code in DAY_CODES:
            display, unit = value, "jours"
        elif code in PERCENT_CODES:
            display, unit = value * 100, "%"
        else:
            display, unit = value, "×"
        ws.cell(row=row_index, column=1, value=row["name"]).border = BORDER
        cell = ws.cell(row=row_index, column=2, value=display)
        cell.border = BORDER
        if isinstance(display, float):
            cell.number_format = "#,##0.000"
        ws.cell(row=row_index, column=3, value=unit).border = BORDER
        row_index += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 78
    ws.column_dimensions["C"].width = 10


def build_report(
    *,
    filing: dict,
    values: list[dict],
    ratios: list[dict],
    checks: list[dict],
    extraction_sha256: str,
    ratios_sha256: str,
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Veda"
    workbook.properties.lastModifiedBy = "Veda"
    workbook.properties.title = f"Rapport de dépôt — {filing['issuer']} {str(filing['period'])[:4]}"
    workbook.properties.created = FIXED_DATE
    workbook.properties.modified = FIXED_DATE

    summary = workbook.active
    summary.title = "Synthèse"
    _summary_sheet(summary, filing, ratios, checks, extraction_sha256, ratios_sha256)

    lines = workbook.create_sheet("Postes")
    line_rows = [
        [
            row["line_item"],
            row["label_raw"],
            row["value"],
            _mtnd(row["value"]),
            row["page"],
            row["confidence"],
        ]
        for row in sorted(values, key=lambda item: item["line_item"])
    ]
    _write_table(
        lines,
        ["Clé", "Libellé", "Valeur (TND)", "Valeur (MTND)", "Page", "Confiance"],
        line_rows,
        ["", "", "#,##0", "#,##0.000", "0", "0.000"],
    )
    _auto_width(lines)

    ratio_sheet = workbook.create_sheet("Ratios")
    ratio_rows = [
        [
            row["code"],
            row["name"],
            row["value"],
            row["year"],
            row["formula_version"],
        ]
        for row in sorted(ratios, key=lambda item: item["code"])
    ]
    _write_table(
        ratio_sheet,
        ["Code", "Libellé", "Valeur", "Exercice", "Version formule"],
        ratio_rows,
        ["", "", "0.000", "0", ""],
    )
    for index, row in enumerate(ratio_rows, start=2):
        ratio_sheet.cell(row=index, column=3).number_format = _ratio_format(row[0])
    _auto_width(ratio_sheet)

    checks_sheet = workbook.create_sheet("Contrôles")
    check_rows = [
        [
            check["id"],
            "Conforme" if check["status"] == "pass" else "Écart",
            check["detail"],
        ]
        for check in checks
    ]
    _write_table(
        checks_sheet,
        ["Contrôle", "Statut", "Détail"],
        check_rows,
        ["", "", ""],
    )
    for index, check in enumerate(checks, start=2):
        cell = checks_sheet.cell(row=index, column=2)
        cell.font = Font(
            bold=True, color="1E7B45" if check["status"] == "pass" else "B03A3A"
        )
    _auto_width(checks_sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_STORED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = CORE_TIMESTAMP.sub(
                    rb"\g<1>2026-01-01T00:00:00Z\g<2>", content
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_STORED
            fixed.external_attr = 0o600 << 16
            fixed.internal_attr = 0
            fixed.create_system = 0
            target.writestr(fixed, content)
    return output.getvalue()


def report_sha256(data: bytes) -> str:
    return sha256_bytes(data)


def _cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def preview_report(data: bytes, max_rows: int = 300, max_cols: int = 30) -> list[dict]:
    """Read the generated workbook back into JSON-friendly sheet rows."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows = []
            truncated = False
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    truncated = True
                    break
                values = [_cell_value(value) for value in row[:max_cols]]
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
            sheets.append(
                {"name": worksheet.title, "rows": rows, "truncated": truncated}
            )
        return sheets
    finally:
        workbook.close()
